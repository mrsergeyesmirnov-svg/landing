"""CI tests use disposable PostgreSQL, never production credentials or customer data."""
import asyncio
import io
import json
import os
import unittest
from pathlib import Path
from unittest.mock import patch

import asyncpg
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

os.environ["SESSION_SECRET"] = "ci-only-secret-at-least-thirty-two-characters"
os.environ["PLATFORM_USERS_JSON"] = '[{"username":"export-test","password_hash":"unused"}]'
import main
import master_app
import master_audit
import audit_exports

BLOCKS = ["Концепция и гость", "Сервис", "Команда", "Менеджмент", "Операции",
          "Кухня", "Продажи", "Финансы", "Маркетинг", "Развитие"]


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        dsn = os.environ.get("TEST_DATABASE_URL", "")
        if not dsn or not any(host in dsn for host in ("@localhost:", "@127.0.0.1:")):
            raise RuntimeError("Tests require TEST_DATABASE_URL pointing to disposable localhost Postgres")
        os.environ["PLATFORM_USERS_JSON"] = '[{"username":"export-test","password_hash":"unused"}]'
        os.environ["DATABASE_URL"] = dsn
        cls.client_context = TestClient(master_app.app, base_url="https://testserver")
        cls.client = cls.client_context.__enter__()
        token, _ = main._make_session("export-test", False)
        cls.client.cookies.set(main.SESSION_COOKIE, token)
        cls.counter = 0

    @classmethod
    def tearDownClass(cls):
        cls.client_context.__exit__(None, None, None)

    def setUp(self):
        main._attempts.clear()

    def create_audit(self, kind="day0", title="Тестовый ресторан «Свет»", client_id=None,
                     complete=True, all_na=False, long=False):
        response = self.client.post("/api/audits", json={
            "restaurant_title": title, "audit_type": kind, "client_id": client_id})
        self.assertEqual(response.status_code, 200, response.text)
        audit = response.json()["audit"]
        for i in range(1, 151):
            section = BLOCKS[(i - 1) // 15]
            payload = dict(item_code=f"{(i-1)//15+1:02}.{(i-1)%15+1:02}",
                           section=section, weight=1 if i % 2 else 2, critical=i % 5 == 0,
                           standard="Стандарт обслуживания гостя: " + section,
                           evidence_hint="Наблюдение смены и проверка действующих регламентов.",
                           score=i % 3, is_na=all_na or i > 135,
                           comment="Подтверждено наблюдением аудитора.")
            if i == 15:
                payload["comment"] = '=HYPERLINK("https://example.invalid","unsafe")'
            if long and i == 30:
                payload["standard"] = ("Очень длинный стандарт <&> " * 100)[:1800]
                payload["comment"] = ("Длинное наблюдение <&> " * 300)[:5000]
                payload["evidence_hint"] = ("Длинное подтверждение " * 100)[:1800]
            saved = self.client.put(f'/api/audits/{audit["id"]}/answers/{i}', json=payload)
            self.assertEqual(saved.status_code, 200, saved.text)
        if complete:
            result = self.client.post(f'/api/audits/{audit["id"]}/complete')
            self.assertEqual(result.status_code, 200, result.text)
            audit = result.json()["audit"]
        return audit

    def export(self, audit, format):
        response = self.client.get(f'/api/audits/{audit["id"]}/export/{format}')
        self.assertEqual(response.status_code, 200, response.text[:300] if response.status_code != 200 else "")
        self.assertEqual(response.headers["cache-control"], "private, no-store")
        self.assertIn("attachment;", response.headers["content-disposition"])
        return response

    def test_end_to_end_snapshot_and_documents(self):
        client = self.client.post("/api/clients", json={"name": "Экспорт CI"}).json()["client"]
        audits = [self.create_audit(kind, client_id=client["id"]) for kind in ("day0", "day30", "day60")]
        audit = audits[-1]
        self.assertEqual(self.client.post(f'/api/audits/{audits[1]["id"]}/baseline').status_code, 200)
        pdf, excel = self.export(audit, "pdf"), self.export(audit, "xlsx")
        self.assertEqual(pdf.headers["x-audit-snapshot"], excel.headers["x-audit-snapshot"])
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertTrue(excel.content.startswith(b"PK"))
        book = load_workbook(io.BytesIO(excel.content), data_only=False)
        self.assertEqual(book["Dashboard"]["B10"].value, audit["overall_score"])
        self.assertEqual(book["Dashboard"]["E10"].value, audit["red_flags_count"])
        self.assertEqual(book["Master Audit"].max_row, 151)
        self.assertEqual(book["Master Audit"]["H16"].data_type, "s")
        self.assertTrue(book["Master Audit"]["H16"].value.startswith("=HYPERLINK"))
        self.assertEqual(book["Master Audit"]["G151"].value, "N/A")
        self.assertEqual(book["Day 0-30-60"]["D2"].value, audit["id"])
        self.assertEqual(book["Healthy Baseline"]["B3"].value, audits[1]["id"])
        self.assertEqual(len(book["Dashboard"]._charts), 1)
        self.assertIsNotNone(book["Master Audit"].auto_filter.ref)
        reader = PdfReader(io.BytesIO(pdf.content))
        text = "\n".join(page.extract_text() for page in reader.pages)
        self.assertIn("Академия счастья", text)
        self.assertIn(f'{audit["overall_score"]:.1f}', text)
        self.assertIn("Healthy Baseline", text)
        self.assertIn("Day 30", text)
        self.assertGreaterEqual(len(reader.pages), 4)
        path = Path("test-artifacts")
        path.mkdir(exist_ok=True)
        (path / "sample-report.pdf").write_bytes(pdf.content)
        (path / "sample-report.xlsx").write_bytes(excel.content)
        (path / "README.txt").write_text("Synthetic CI data only; not a customer report.\n", encoding="utf-8")

    def test_title_isolation_no_future_and_na(self):
        first = self.create_audit(title="Уникальный ресторан без CRM", all_na=True)
        unrelated = self.create_audit("day30", title="ДРУГОЙ ресторан без CRM")
        future = self.create_audit("day60", title="Уникальный ресторан без CRM")
        book = load_workbook(io.BytesIO(self.export(first, "xlsx").content))
        self.assertIsNone(book["Day 0-30-60"]["C2"].value)
        self.assertIsNone(book["Day 0-30-60"]["D2"].value)
        self.assertEqual(book["Dashboard"]["B10"].value, 0)
        self.assertNotIn("Healthy Baseline", book.sheetnames)
        pdf = self.export(first, "pdf").content
        text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
        self.assertIn("N/A: 150", text)
        self.assertNotIn("ДРУГОЙ", text)

    def test_route_guards_and_existing_ai(self):
        saved_cookies = dict(self.client.cookies)
        self.client.cookies.clear()
        try:
            self.assertEqual(self.client.get("/api/audits/missing/export/pdf").status_code, 401)
        finally:
            self.client.cookies.update(saved_cookies)
        self.assertIn("/api/audits/{audit_id}/ai-ingest",
                      {route.path for route in master_app.app.routes})

    def test_long_text_and_escaping(self):
        audit = self.create_audit(title="Ресторан <&> " * 12, long=True)
        pdf = self.export(audit, "pdf")
        self.assertGreater(len(PdfReader(io.BytesIO(pdf.content)).pages), 4)
        book = load_workbook(io.BytesIO(self.export(audit, "xlsx").content))
        self.assertEqual(len(book["Master Audit"]["H31"].value), 5000)

    def test_missing_draft_format_and_mismatch(self):
        self.assertEqual(self.client.get("/api/audits/missing/export/pdf").status_code, 404)
        self.assertEqual(self.client.get("/api/audits/missing/export/csv").status_code, 400)
        draft = self.create_audit(complete=False)
        self.assertEqual(self.client.get(f'/api/audits/{draft["id"]}/export/pdf').status_code, 409)
        self.client.post(f'/api/audits/{draft["id"]}/complete')
        async def corrupt():
            conn = await asyncpg.connect(os.environ["TEST_DATABASE_URL"])
            try:
                await conn.execute("UPDATE academy_audit_sessions SET overall_score=99.9 WHERE id=$1", draft["id"])
            finally:
                await conn.close()
        asyncio.run(corrupt())
        self.assertEqual(self.client.get(f'/api/audits/{draft["id"]}/export/xlsx').status_code, 409)


if __name__ == "__main__":
    unittest.main()
