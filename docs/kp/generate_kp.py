#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Внутренние КП — PDF + DOCX. Не выкладывать на публичный лендинг."""

from pathlib import Path

from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
)

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path(__file__).resolve().parent

pdfmetrics.registerFont(TTFont("Serif", "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"))
pdfmetrics.registerFont(TTFont("SerifB", "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf"))
pdfmetrics.registerFont(TTFont("Sans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("SansB", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))

INK = HexColor("#1A1612")
MUTED = HexColor("#5C5348")
GOLD = HexColor("#D4A84B")
GOLD_INK = HexColor("#B8892E")
CREAM = HexColor("#FFFCF7")
LINE = HexColor("#EFE6D6")
WARM = HexColor("#FBF3E4")
ROW2 = HexColor("#FFF9F0")

CONTACT = "Telegram: @yank0vski · pulseteam.online"
BRAND = "Академия счастья · «Состояние смены»"


def S():
    return {
        "brand": ParagraphStyle("brand", fontName="Serif", fontSize=11, textColor=GOLD_INK, leading=14, spaceAfter=4),
        "kicker": ParagraphStyle("kicker", fontName="SansB", fontSize=8, textColor=GOLD_INK, leading=11, spaceAfter=8),
        "h1": ParagraphStyle("h1", fontName="SerifB", fontSize=20, textColor=INK, leading=26, spaceAfter=10),
        "h2": ParagraphStyle("h2", fontName="SerifB", fontSize=13, textColor=INK, leading=17, spaceBefore=14, spaceAfter=7),
        "lead": ParagraphStyle("lead", fontName="Sans", fontSize=10.5, textColor=MUTED, leading=15, alignment=TA_JUSTIFY, spaceAfter=10),
        "body": ParagraphStyle("body", fontName="Sans", fontSize=10, textColor=INK, leading=14, alignment=TA_JUSTIFY, spaceAfter=6),
        "li": ParagraphStyle("li", fontName="Sans", fontSize=10, textColor=INK, leading=14, spaceAfter=3),
        "price": ParagraphStyle("price", fontName="SerifB", fontSize=18, textColor=GOLD_INK, leading=22, spaceBefore=2, spaceAfter=6),
        "note": ParagraphStyle("note", fontName="Sans", fontSize=8.5, textColor=MUTED, leading=12, spaceBefore=6),
        "cell": ParagraphStyle("cell", fontName="Sans", fontSize=9.5, textColor=INK, leading=13, alignment=TA_CENTER),
        "cellL": ParagraphStyle("cellL", fontName="Sans", fontSize=9.5, textColor=INK, leading=13),
        "cellB": ParagraphStyle("cellB", fontName="SansB", fontSize=9.5, textColor=INK, leading=13, alignment=TA_CENTER),
        "cellBL": ParagraphStyle("cellBL", fontName="SansB", fontSize=9.5, textColor=INK, leading=13),
    }


def hr():
    return HRFlowable(width="100%", thickness=0.6, color=LINE, spaceBefore=2, spaceAfter=8)


def bullets(items, st):
    return [Paragraph("•  " + t, st["li"]) for t in items]


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(CREAM)
    canvas.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
    canvas.setStrokeColor(GOLD)
    canvas.setLineWidth(2)
    canvas.line(18 * mm, A4[1] - 12 * mm, A4[0] - 18 * mm, A4[1] - 12 * mm)
    canvas.setFont("Sans", 7.5)
    canvas.setFillColor(MUTED)
    canvas.drawString(18 * mm, 12 * mm, BRAND)
    canvas.drawRightString(A4[0] - 18 * mm, 12 * mm, CONTACT)
    canvas.restoreState()


def build_pdf(path, title, blocks):
    st = S()
    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
        title=title, author="Академия счастья",
    )
    story = [
        Paragraph("АКАДЕМИЯ СЧАСТЬЯ", st["brand"]),
        Paragraph("КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ", st["kicker"]),
        Paragraph(title, st["h1"]),
        hr(),
    ]
    story.extend(blocks)
    story += [
        Spacer(1, 12),
        hr(),
        Paragraph(
            "Внутренний документ для переговоров. Итоговые условия фиксируются в договоре / счёте.",
            st["note"],
        ),
    ]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def run_font(run, size=10, bold=False, color=None, serif=False):
    run.font.name = "DejaVu Serif" if serif else "DejaVu Sans"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), run.font.name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def shade(cell, fill="FBF3E4"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shd)


def build_docx(path, title, sections):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(1.8)
        sec.bottom_margin = Cm(1.8)
        sec.left_margin = Cm(2)
        sec.right_margin = Cm(2)

    p = doc.add_paragraph()
    r = p.add_run("АКАДЕМИЯ СЧАСТЬЯ")
    run_font(r, 11, False, RGBColor(0xB8, 0x89, 0x2E), True)

    p = doc.add_paragraph()
    r = p.add_run("КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ")
    run_font(r, 8, True, RGBColor(0xB8, 0x89, 0x2E))

    p = doc.add_paragraph()
    r = p.add_run(title)
    run_font(r, 18, True, RGBColor(0x1A, 0x16, 0x12), True)

    p = doc.add_paragraph()
    r = p.add_run(BRAND + " · " + CONTACT)
    run_font(r, 9, False, RGBColor(0x5C, 0x53, 0x48))

    for sec in sections:
        t = sec["type"]
        if t == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            r = p.add_run(sec["text"])
            run_font(r, 13, True, RGBColor(0x1A, 0x16, 0x12), True)
        elif t == "p":
            p = doc.add_paragraph()
            r = p.add_run(sec["text"])
            col = RGBColor(0x5C, 0x53, 0x48) if sec.get("muted") else RGBColor(0x1A, 0x16, 0x12)
            run_font(r, 10, False, col)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.35
        elif t == "price":
            p = doc.add_paragraph()
            r = p.add_run(sec["text"])
            run_font(r, 16, True, RGBColor(0xB8, 0x89, 0x2E), True)
        elif t == "bullets":
            for item in sec["items"]:
                p = doc.add_paragraph()
                r = p.add_run("•  " + item)
                run_font(r, 10, False, RGBColor(0x1A, 0x16, 0x12))
                p.paragraph_format.left_indent = Cm(0.25)
                p.paragraph_format.space_after = Pt(3)
        elif t == "table":
            rows = sec["rows"]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    cell = table.cell(i, j)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    r = p.add_run(val)
                    run_font(r, 9, bold=(i == 0), color=RGBColor(0x1A, 0x16, 0x12))
                    if i == 0:
                        shade(cell)
            doc.add_paragraph()

    p = doc.add_paragraph()
    r = p.add_run("Внутренний документ для переговоров. Итоговые условия фиксируются в договоре / счёте.")
    run_font(r, 8, False, RGBColor(0x5C, 0x53, 0x48))
    doc.save(str(path))


def make_tariff_table(st):
    data = [
        [Paragraph("Посадочные места в точке", st["cellBL"]), Paragraph("Подписка / мес", st["cellB"])],
        [Paragraph("До 40 мест", st["cellL"]), Paragraph("<b>2 990 ₽</b>", st["cell"])],
        [Paragraph("41–100 мест", st["cellL"]), Paragraph("<b>6 990 ₽</b>", st["cell"])],
        [Paragraph("101–160 мест", st["cellL"]), Paragraph("<b>8 990 ₽</b>", st["cell"])],
        [Paragraph("Более 160 мест / сеть", st["cellL"]), Paragraph("<b>индивидуально</b>", st["cell"])],
    ]
    t = Table(data, colWidths=[100 * mm, 55 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), WARM),
        ("GRID", (0, 0), (-1, -1), 0.45, LINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [CREAM, ROW2]),
    ]))
    return t


def kp1():
    st = S()
    title = "Подписка на Telegram-бота «Состояние смены»"
    pdf = [
        Paragraph(
            "Ежедневный мониторинг состояния команды в Telegram: короткие опросы после смены, "
            "сводки управляющему, «горящие вопросы» и ИИ-помощник. Вы видите напряжение в зале и на кухне "
            "раньше, чем оно станет увольнениями, жалобами гостей или просадкой выручки.",
            st["lead"],
        ),
        Paragraph("Что входит в подписку", st["h2"]),
        *bullets([
            "Ежедневные check-in сотрудников (зал / кухня) — 10–30 секунд на человека",
            "Сводка управляющему: настроение смены, повторяющиеся темы, риски",
            "«Горящие вопросы» — сигналы, которые нельзя откладывать",
            "ИИ-помощник: формулировки, подсказки, разбор паттернов",
            "Разделение контуров зал / кухня и отчёты по точке",
            "Подключение и онбординг команды на старте пилота",
        ], st),
        Paragraph("Тарифы — по посадочным местам", st["h2"]),
        Paragraph(
            "Метрика тарифа — вместимость зала (посадочные места в точке), а не число отметившихся сотрудников. "
            "Так проще планировать бюджет: цена привязана к масштабу ресторана. Суммы тарифов сохранены.",
            st["body"],
        ),
        Spacer(1, 6),
        make_tariff_table(st),
        Spacer(1, 8),
        Paragraph("Пилот 1 месяц — бесплатно", st["h2"]),
        Paragraph(
            "Полный контур бота без оплаты в первый месяц: опросы, сводки, горящие вопросы, ИИ-помощник. "
            "Без обязательств продлевать подписку. После пилота — переход на тариф по посадочным местам или пауза без штрафов.",
            st["body"],
        ),
        Paragraph("Как запускаем", st["h2"]),
        *bullets([
            "Короткий созвон: точка, посадочные места, роли (управляющий / шеф / HR)",
            "Подключение бота и инструкции для команды",
            "Неделя 1: привычка отвечать + первая сводка",
            "Недели 2–4: рабочие сигналы и разбор с вами",
            "Итог пилота: что увидели, что менять, решение по подписке",
        ], st),
        Paragraph("Для кого", st["h2"]),
        *bullets([
            "Собственник или управляющий, которому нужна картина по команде без «мне кажется»",
            "Точки и сети, где важно ловить выгорание и сбои до гостевых жалоб",
            "HR и операционный контур с запросом на регулярный пульс смены",
        ], st),
        Paragraph("Следующий шаг", st["h2"]),
        Paragraph(
            "Напишите в Telegram @yank0vski — укажите город, концепцию и число посадочных мест. Подключим пилот и пришлём доступ.",
            st["body"],
        ),
    ]
    docx = [
        {"type": "p", "muted": True, "text":
            "Ежедневный мониторинг состояния команды в Telegram: опросы после смены, сводки управляющему, "
            "«горящие вопросы» и ИИ-помощник. Напряжение видно раньше увольнений и жалоб гостей."},
        {"type": "h2", "text": "Что входит в подписку"},
        {"type": "bullets", "items": [
            "Ежедневные check-in (зал / кухня) — 10–30 секунд",
            "Сводка управляющему: настроение, темы, риски",
            "«Горящие вопросы»",
            "ИИ-помощник",
            "Контуры зал / кухня и отчёты по точке",
            "Онбординг команды на старте пилота",
        ]},
        {"type": "h2", "text": "Тарифы — по посадочным местам"},
        {"type": "p", "text": "Метрика — посадочные места в точке, не число отметившихся. Цены сохранены."},
        {"type": "table", "rows": [
            ["Посадочные места в точке", "Подписка / мес"],
            ["До 40 мест", "2 990 ₽"],
            ["41–100 мест", "6 990 ₽"],
            ["101–160 мест", "8 990 ₽"],
            ["Более 160 мест / сеть", "индивидуально"],
        ]},
        {"type": "h2", "text": "Пилот 1 месяц — бесплатно"},
        {"type": "p", "text": "Полный контур бота без оплаты в первый месяц. Без обязательств продлевать подписку."},
        {"type": "h2", "text": "Как запускаем"},
        {"type": "bullets", "items": [
            "Созвон: точка, посадочные места, роли",
            "Подключение бота и инструкции",
            "Неделя 1: привычка + первая сводка",
            "Недели 2–4: сигналы и разбор",
            "Итог пилота: решение по подписке",
        ]},
        {"type": "h2", "text": "Следующий шаг"},
        {"type": "p", "text": "Telegram @yank0vski — город, концепция, посадочные места. Подключим пилот."},
    ]
    return title, pdf, docx


def kp2():
    st = S()
    title = "Аудит «Состояние смены» — диагностика команды и операций"
    pdf = [
        Paragraph(
            "Экспресс-аудит даёт ясную картину: что происходит с людьми, гостем, процессами и результатом — "
            "без презентации на 40 слайдов. Вы получаете карту причин, приоритеты и понятный следующий шаг.",
            st["lead"],
        ),
        Paragraph("20 000 ₽", st["price"]),
        Paragraph(
            "Фиксированная стоимость экспресс-аудита для одной точки (ориентир). "
            "Проектная работа большего масштаба считается отдельно после аудита.",
            st["body"],
        ),
        Paragraph("Что входит", st["h2"]),
        *bullets([
            "Установочная сессия с собственником / управляющим (цели, боли, ограничения)",
            "Чек-листы диагностики по четырём системам: люди, гости, финансы, процессы",
            "Сбор сигналов: интервью / наблюдение / данные смены (по договорённости)",
            "Разбор «симптом → причина»: где теряется энергия команды и сервиса",
            "Письменный итог: карта состояния, риски, быстрые победы, приоритеты",
            "Рекомендация формата продолжения: бот / проект внедрения / пауза",
        ], st),
        Paragraph("Чек-листы (рамка аудита)", st["h2"]),
        *bullets([
            "Люди: нагрузка, опора на ключевых, слышимость команды, удержание",
            "Гости: сервисный контур, жалобы, возвраты, качество опыта",
            "Финансы: где состояние команды бьёт по выручке и списаниям",
            "Процессы: смена, стандарты, эскалация, «кто узнаёт первым»",
        ], st),
        Paragraph("Установочные сессии", st["h2"]),
        *bullets([
            "Старт: ожидания собственника, что считается успехом аудита",
            "Промежуточная сверка (по необходимости): уточнение гипотез",
            "Финал: разбор результатов и развилка по следующим этапам",
        ], st),
        Paragraph("Дальнейшие этапы после аудита", st["h2"]),
        *bullets([
            "Этап A — Мониторинг: бот «Состояние смены» (пилот 1 месяц бесплатно)",
            "Этап B — Проект внедрения: план изменений, работа с управляющими, закрепление привычек",
            "Этап C — Контроль «было → стало»: повторные замеры, корректировка, отчёт собственнику",
            "Этап D — Сопровождение: регулярные сессии + аналитика бота",
        ], st),
        Paragraph("Как проходит", st["h2"]),
        *bullets([
            "День 0: бриф и контакты",
            "1–5 дней: диагностика по чек-листам и сигналам",
            "Итоговая сессия: выводы и решение о следующем этапе",
            "Документ аудита остаётся у вас",
        ], st),
        Paragraph("Результат для собственника", st["h2"]),
        *bullets([
            "Понимание, где ресторан теряет энергию — не на уровне лозунгов",
            "Приоритеты: что делать сначала, а что не трогать сейчас",
            "Ясный выбор: только бот / проект / ничего не менять осознанно",
        ], st),
        Paragraph("Следующий шаг", st["h2"]),
        Paragraph(
            "Напишите @yank0vski — коротко опишите точку и что беспокоит. "
            "Согласуем дату установочной сессии и старт аудита.",
            st["body"],
        ),
    ]
    docx = [
        {"type": "p", "muted": True, "text":
            "Экспресс-аудит: люди, гость, процессы, результат — без презентации на 40 слайдов."},
        {"type": "price", "text": "20 000 ₽"},
        {"type": "p", "text":
            "Фиксированная стоимость экспресс-аудита для одной точки (ориентир). "
            "Крупный проект — отдельно после аудита."},
        {"type": "h2", "text": "Что входит"},
        {"type": "bullets", "items": [
            "Установочная сессия с собственником / управляющим",
            "Чек-листы: люди, гости, финансы, процессы",
            "Сбор сигналов: интервью / наблюдение / данные смены",
            "Разбор «симптом → причина»",
            "Письменный итог: карта, риски, быстрые победы, приоритеты",
            "Рекомендация формата продолжения",
        ]},
        {"type": "h2", "text": "Чек-листы"},
        {"type": "bullets", "items": [
            "Люди: нагрузка, опора, слышимость, удержание",
            "Гости: сервис, жалобы, возвраты, опыт",
            "Финансы: влияние состояния команды на результат",
            "Процессы: смена, стандарты, эскалация",
        ]},
        {"type": "h2", "text": "Установочные сессии"},
        {"type": "bullets", "items": [
            "Старт: ожидания и критерий успеха",
            "Промежуточная сверка (по необходимости)",
            "Финал: выводы и развилка этапов",
        ]},
        {"type": "h2", "text": "Дальнейшие этапы"},
        {"type": "bullets", "items": [
            "A — Мониторинг ботом (пилот 1 месяц бесплатно)",
            "B — Проект внедрения изменений",
            "C — Контроль «было → стало»",
            "D — Сопровождение + аналитика",
        ]},
        {"type": "h2", "text": "Следующий шаг"},
        {"type": "p", "text": "Telegram @yank0vski — опишите точку и запрос. Согласуем установочную сессию."},
    ]
    return title, pdf, docx


def kp3():
    st = S()
    title = "Бесплатная установочная сессия с собственником (2–4 часа)"
    pdf = [
        Paragraph(
            "Первый разговор без продажи «в лоб»: разбираем, что происходит в ресторане глазами собственника, "
            "где болит команда и сервис, и какие есть реалистичные пути дальше. "
            "Формат — 2–4 часа (очно или онлайн).",
            st["lead"],
        ),
        Paragraph("0 ₽", st["price"]),
        Paragraph(
            "Сессия бесплатная. Это вход в диалог, а не обязательство покупать аудит или подписку.",
            st["body"],
        ),
        Paragraph("Что получит собственник", st["h2"]),
        *bullets([
            "Структурированное поле проблем: люди / гости / деньги / процессы — без хаоса в голове",
            "Отделение симптомов от причин (что «лечили» уже не раз и почему возвращается)",
            "Честная оценка: что можно увидеть самим, а где нужна диагностика извне",
            "Короткий список быстрых действий на 7–14 дней (если уместно)",
            "Понятная карта вариантов сотрудничества — без давления",
        ], st),
        Paragraph("Что узнает нового", st["h2"]),
        *bullets([
            "Как состояние команды проявляется в смене раньше цифр и отзывов",
            "Где управляющий перегружен «героизмом» вместо системы",
            "Какие сигналы можно собирать ежедневно (и зачем это собственнику)",
            "Чем аудит отличается от «ещё одного тренинга»",
            "Как пилот бота даёт точку отсчёта до любых изменений",
        ], st),
        Paragraph("Как проходит сессия (2–4 часа)", st["h2"]),
        *bullets([
            "Контекст: точка, команда, роль собственника в операционке",
            "Разбор текущих болей и повторяющихся сбоев",
            "Проявление четырёх систем ресторана на вашем примере",
            "Развилка: что делать дальше и чего не делать сейчас",
            "Фиксация договорённостей и краткий конспект",
        ], st),
        Paragraph("Пути развития сотрудничества", st["h2"]),
        *bullets([
            "1) Ничего не покупать — уходите с ясностью и коротким планом на две недели",
            "2) Пилот бота «Состояние смены» — 1 месяц бесплатно, далее тариф по посадочным местам",
            "3) Экспресс-аудит — 20 000 ₽: чек-листы, установочные сессии, письменный итог",
            "4) Проект внедрения после аудита — изменения + контроль «было → стало»",
            "5) Сопровождение — регулярные сессии для собственника / управляющих",
        ], st),
        Paragraph("Кому особенно полезно", st["h2"]),
        *bullets([
            "Собственник чувствует «что-то не так», но нет ясной картины",
            "Управляющие выгорают, проблемы циклично возвращаются",
            "Нужно понять: бот, аудит или достаточно внутренней дисциплины",
        ], st),
        Paragraph("Следующий шаг", st["h2"]),
        Paragraph(
            "Напишите @yank0vski: имя, город, концепция, удобный формат (очно / Zoom) и окно по датам. "
            "Забронируем 2–4 часа.",
            st["body"],
        ),
    ]
    docx = [
        {"type": "p", "muted": True, "text":
            "Первый разговор без продажи «в лоб»: что происходит в ресторане, где болит команда и сервис, "
            "какие пути дальше. Формат — 2–4 часа (очно или онлайн)."},
        {"type": "price", "text": "0 ₽"},
        {"type": "p", "text": "Сессия бесплатная. Без обязательства покупать аудит или подписку."},
        {"type": "h2", "text": "Что получит собственник"},
        {"type": "bullets", "items": [
            "Структура проблем: люди / гости / деньги / процессы",
            "Симптомы vs причины",
            "Оценка: что видно самим, где нужна диагностика",
            "Быстрые действия на 7–14 дней (если уместно)",
            "Карта вариантов сотрудничества без давления",
        ]},
        {"type": "h2", "text": "Что узнает нового"},
        {"type": "bullets", "items": [
            "Как состояние команды проявляется раньше цифр и отзывов",
            "Где «героизм» заменяет систему",
            "Какие сигналы можно собирать ежедневно",
            "Чем аудит отличается от тренинга",
            "Зачем пилот бота как точка отсчёта",
        ]},
        {"type": "h2", "text": "Как проходит (2–4 часа)"},
        {"type": "bullets", "items": [
            "Контекст точки и роли собственника",
            "Разбор болей и повторяющихся сбоев",
            "Четыре системы на вашем примере",
            "Развилка следующих шагов",
            "Краткий конспект договорённостей",
        ]},
        {"type": "h2", "text": "Пути сотрудничества"},
        {"type": "bullets", "items": [
            "1) Уйти с ясностью и коротким планом",
            "2) Пилот бота — 1 месяц бесплатно",
            "3) Экспресс-аудит — 20 000 ₽",
            "4) Проект внедрения после аудита",
            "5) Регулярное сопровождение",
        ]},
        {"type": "h2", "text": "Следующий шаг"},
        {"type": "p", "text":
            "Telegram @yank0vski — имя, город, концепция, очно/Zoom, даты. Забронируем 2–4 часа."},
    ]
    return title, pdf, docx


def main():
    jobs = [
        ("kp-01-bot-podpiska", kp1),
        ("kp-02-audit", kp2),
        ("kp-03-ustanovochnaya-sessiya", kp3),
    ]
    for slug, fn in jobs:
        title, pdf_blocks, docx_sections = fn()
        build_pdf(OUT / f"{slug}.pdf", title, pdf_blocks)
        build_docx(OUT / f"{slug}.docx", title, docx_sections)
        print("OK", slug)


if __name__ == "__main__":
    main()
